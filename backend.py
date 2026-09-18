"""
backend.py — Logique métier de l'assistant AM TRANSIT (sans dépendance à Streamlit).

Contient :
- La bibliothèque de documents persistante (storage/) : upload, listing, chargement en DataFrame
- L'outil "query_dataframe" (exécution sécurisée d'expressions pandas)
- L'extraction de texte PDF
- L'historique des conversations (SQLite)
- Les exports (Excel / PDF) d'une conversation

Ce module ne dépend d'aucune UI et peut être testé/importé indépendamment.
"""
from __future__ import annotations

import os
import re
import json
import sqlite3
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd

BASE_DIR = Path(__file__).resolve().parent
STORAGE_DIR = BASE_DIR / "storage"
DB_PATH = STORAGE_DIR / "chat_history.db"

STORAGE_DIR.mkdir(parents=True, exist_ok=True)

ALLOWED_EXTENSIONS = {".xlsx", ".xls", ".csv", ".pdf"}


# --------------------------------------------------------------------------
# Bibliothèque de documents
# --------------------------------------------------------------------------

def save_uploaded_file(filename: str, content: bytes) -> Path:
    """Enregistre un fichier uploadé dans la bibliothèque persistante storage/.

    Si un fichier du même nom existe déjà, il est remplacé (mise à jour).
    """
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError(
            f"Type de fichier non supporté : {ext}. "
            f"Formats acceptés : {', '.join(sorted(ALLOWED_EXTENSIONS))}"
        )
    safe_name = re.sub(r"[^A-Za-z0-9._\-]+", "_", filename)
    dest = STORAGE_DIR / safe_name
    dest.write_bytes(content)
    return dest


def list_documents() -> list[dict[str, Any]]:
    """Liste tous les documents de la bibliothèque avec un résumé de leur structure."""
    docs = []
    for path in sorted(STORAGE_DIR.iterdir()):
        if path.is_dir() or path.suffix.lower() not in ALLOWED_EXTENSIONS:
            continue
        entry: dict[str, Any] = {
            "filename": path.name,
            "type": path.suffix.lower().lstrip("."),
            "size_kb": round(path.stat().st_size / 1024, 1),
        }
        try:
            if path.suffix.lower() in (".xlsx", ".xls"):
                xl = pd.ExcelFile(path)
                entry["sheets"] = {
                    sheet: list(pd.read_excel(path, sheet_name=sheet, nrows=0).columns)
                    for sheet in xl.sheet_names
                }
            elif path.suffix.lower() == ".csv":
                entry["sheets"] = {"csv": list(pd.read_csv(path, nrows=0).columns)}
            elif path.suffix.lower() == ".pdf":
                entry["sheets"] = None
        except Exception as exc:  # noqa: BLE001
            entry["error"] = str(exc)
        docs.append(entry)
    return docs


def remove_document(filename: str) -> bool:
    safe_name = re.sub(r"[^A-Za-z0-9._\-]+", "_", filename)
    path = STORAGE_DIR / safe_name
    if path.exists():
        path.unlink()
        return True
    return False


_df_cache: dict[str, pd.DataFrame] = {}


def _load_dataframe(filename: str, sheet: str | None) -> pd.DataFrame:
    key = f"{filename}::{sheet}"
    if key in _df_cache:
        return _df_cache[key]
    path = STORAGE_DIR / re.sub(r"[^A-Za-z0-9._\-]+", "_", filename)
    if not path.exists():
        raise FileNotFoundError(f"Document introuvable dans la bibliothèque : {filename}")
    if path.suffix.lower() in (".xlsx", ".xls"):
        df = pd.read_excel(path, sheet_name=sheet or 0)
    elif path.suffix.lower() == ".csv":
        df = pd.read_csv(path)
    else:
        raise ValueError(f"Ce type de fichier n'est pas une table : {filename}")
    _df_cache[key] = df
    return df


_FORBIDDEN_TOKENS = (
    "__", "import", "open(", "exec(", "eval(", "os.", "sys.", "subprocess",
    "globals", "locals", "compile(", "input(", "breakpoint",
)


def query_dataframe(filename: str, expression: str, sheet: str | None = None, max_rows: int = 50) -> str:
    """Évalue une expression pandas sur un document de la bibliothèque, de façon restreinte.

    L'expression a accès à `df` (le DataFrame chargé) et `pd`. Exemples :
        "df.groupby('N° compte général')[['Débit','Crédit']].sum()"
        "df[df['Code journal'] == 'ACH001'].head(20)"
    """
    lowered = expression.lower()
    if any(tok in lowered for tok in _FORBIDDEN_TOKENS):
        return "Erreur : expression refusée (mot-clé interdit)."

    try:
        df = _load_dataframe(filename, sheet)
    except Exception as exc:  # noqa: BLE001
        return f"Erreur de chargement du document : {exc}"

    safe_globals = {"__builtins__": {}}
    safe_locals = {"df": df, "pd": pd}
    try:
        result = eval(expression, safe_globals, safe_locals)  # noqa: S307
    except Exception as exc:  # noqa: BLE001
        return f"Erreur d'exécution de l'expression : {exc}"

    if isinstance(result, pd.DataFrame):
        truncated = len(result) > max_rows
        out = result.head(max_rows).to_string()
        if truncated:
            out += f"\n… ({len(result)} lignes au total, {max_rows} affichées)"
        return out
    if isinstance(result, pd.Series):
        truncated = len(result) > max_rows
        out = result.head(max_rows).to_string()
        if truncated:
            out += f"\n… ({len(result)} lignes au total, {max_rows} affichées)"
        return out
    return str(result)


def read_document_text(filename: str, max_chars: int = 12000) -> str:
    """Extrait le texte d'un document PDF de la bibliothèque (tronqué à max_chars)."""
    path = STORAGE_DIR / re.sub(r"[^A-Za-z0-9._\-]+", "_", filename)
    if not path.exists():
        return f"Erreur : document introuvable ({filename})"
    if path.suffix.lower() != ".pdf":
        return "Erreur : cet outil ne lit que les PDF. Utilisez query_dataframe pour les fichiers Excel/CSV."
    try:
        from pypdf import PdfReader
    except ImportError:
        return "Erreur : la librairie pypdf n'est pas installée."
    reader = PdfReader(str(path))
    text_parts = []
    for page in reader.pages:
        text_parts.append(page.extract_text() or "")
    full_text = "\n".join(text_parts)
    if len(full_text) > max_chars:
        return full_text[:max_chars] + f"\n… (texte tronqué, {len(full_text)} caractères au total)"
    return full_text


# --------------------------------------------------------------------------
# Historique des conversations (SQLite)
# --------------------------------------------------------------------------

def _get_conn() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS conversations (
            id TEXT PRIMARY KEY,
            title TEXT,
            created_at TEXT
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS messages (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            conversation_id TEXT,
            role TEXT,
            content TEXT,
            created_at TEXT,
            FOREIGN KEY (conversation_id) REFERENCES conversations(id)
        )
        """
    )
    return conn


def create_conversation(title: str = "Nouvelle conversation") -> str:
    conv_id = str(uuid.uuid4())
    with _get_conn() as conn:
        conn.execute(
            "INSERT INTO conversations (id, title, created_at) VALUES (?, ?, ?)",
            (conv_id, title, datetime.utcnow().isoformat()),
        )
    return conv_id


def list_conversations() -> list[dict[str, Any]]:
    with _get_conn() as conn:
        rows = conn.execute(
            "SELECT id, title, created_at FROM conversations ORDER BY created_at DESC"
        ).fetchall()
    return [{"id": r[0], "title": r[1], "created_at": r[2]} for r in rows]


def add_message(conversation_id: str, role: str, content: str) -> None:
    with _get_conn() as conn:
        conn.execute(
            "INSERT INTO messages (conversation_id, role, content, created_at) VALUES (?, ?, ?, ?)",
            (conversation_id, role, content, datetime.utcnow().isoformat()),
        )


def get_messages(conversation_id: str) -> list[dict[str, Any]]:
    with _get_conn() as conn:
        rows = conn.execute(
            "SELECT role, content, created_at FROM messages WHERE conversation_id = ? ORDER BY id ASC",
            (conversation_id,),
        ).fetchall()
    return [{"role": r[0], "content": r[1], "created_at": r[2]} for r in rows]


def rename_conversation(conversation_id: str, title: str) -> None:
    with _get_conn() as conn:
        conn.execute("UPDATE conversations SET title = ? WHERE id = ?", (title, conversation_id))


def delete_conversation(conversation_id: str) -> None:
    with _get_conn() as conn:
        conn.execute("DELETE FROM messages WHERE conversation_id = ?", (conversation_id,))
        conn.execute("DELETE FROM conversations WHERE id = ?", (conversation_id,))


# --------------------------------------------------------------------------
# Export d'une conversation
# --------------------------------------------------------------------------

def export_conversation_xlsx(conversation_id: str, out_path: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill

    messages = get_messages(conversation_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Conversation"
    headers = ["Horodatage", "Rôle", "Message"]
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = header_font
        c.fill = header_fill
    for i, msg in enumerate(messages, start=2):
        ws.cell(row=i, column=1, value=msg["created_at"]).font = Font(name="Arial", size=10)
        ws.cell(row=i, column=2, value=msg["role"]).font = Font(name="Arial", size=10)
        cell = ws.cell(row=i, column=3, value=msg["content"])
        cell.font = Font(name="Arial", size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 100
    out_path = Path(out_path)
    wb.save(out_path)
    return out_path


def export_conversation_pdf(conversation_id: str, out_path: str | Path) -> Path:
    from fpdf import FPDF

    messages = get_messages(conversation_id)
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "Conversation - Assistant AM TRANSIT", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    for msg in messages:
        pdf.set_font("Helvetica", "B", 10)
        label = "Vous" if msg["role"] == "user" else "Assistant"
        pdf.multi_cell(0, 6, f"[{msg['created_at']}] {label} :", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 10)
        safe_content = msg["content"].encode("latin-1", "replace").decode("latin-1")
        pdf.multi_cell(0, 6, safe_content, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)
    out_path = Path(out_path)
    pdf.output(str(out_path))
    return out_path
